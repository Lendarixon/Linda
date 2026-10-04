import io
import json
import subprocess
import sys
import time

import pytest
from linda_desktop import documents


@pytest.fixture(autouse=True)
def isolated_document_spool(monkeypatch,tmp_path):
    monkeypatch.setenv('LINDA_HOME',str(tmp_path/'document-profile'))


def test_docx_real_worker_unicode():
    import docx
    doc = docx.Document()
    doc.add_paragraph('Тестowy документ — Unicode.')
    stream = io.BytesIO()
    doc.save(stream)
    assert documents.extract('test.docx',stream.getvalue())=='Тестowy документ — Unicode.'


def test_pdf_real_worker_and_empty_pdf():
    from fpdf import FPDF
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font('Helvetica',size=12)
    pdf.cell(text='Isolated PDF test')
    assert 'Isolated PDF test' in documents.extract('test.pdf',bytes(pdf.output()))
    from pypdf import PdfWriter
    writer = PdfWriter()
    writer.add_blank_page(width=72,height=72)
    stream = io.BytesIO()
    writer.write(stream)
    assert documents.extract('empty.pdf',stream.getvalue())==''


def test_timeout_kills_owned_worker_and_releases_slot(monkeypatch):
    monkeypatch.setattr(documents,'TIMEOUT',.15)
    from linda_desktop import document_sandbox
    real = document_sandbox.stage_runtime
    def stage(directory):
        command = real(directory)
        from pathlib import Path
        Path(command[1]).write_text('import time; time.sleep(30)',encoding='utf-8')
        return command
    monkeypatch.setattr(document_sandbox,'stage_runtime',stage)
    start = time.monotonic()
    with pytest.raises(ValueError,match='timed out'):
        documents.extract('slow.pdf',b'test')
    assert time.monotonic()-start<5
    assert documents._slots.acquire(blocking=False)
    documents._slots.release()


def test_worker_errors_do_not_leak_paths():
    with pytest.raises(ValueError):
        documents.extract('bad.docx',b'not a zip')
    assert documents.extract('ok.txt','hello'.encode())=='hello'


def test_oversized_pdf_and_output_are_rejected(monkeypatch):
    from linda_desktop import document_worker as worker
    from pypdf import PdfWriter
    monkeypatch.setattr(worker,'MAX_PAGES',1)
    writer = PdfWriter()
    for _ in range(2):
        writer.add_blank_page(width=72,height=72)
    stream = io.BytesIO()
    writer.write(stream)
    with pytest.raises(ValueError,match='page limit'):
        worker.parse('test.pdf',stream.getvalue())
    monkeypatch.setattr(documents,'MAX_INPUT',4)
    with pytest.raises(ValueError,match='input limit'):
        documents.extract('large.pdf',b'12345')


def test_windowed_frozen_command_has_worker_dispatch(monkeypatch,tmp_path):
    monkeypatch.setattr(sys,'frozen',True,raising=False)
    assert documents._command('a.pdf',tmp_path/'in',tmp_path/'out')[1]=='--document-worker'
    # Stable early dispatch must avoid importing the app or configuring profiles.
    from pathlib import Path
    root = Path(__file__).resolve().parents[1]
    entry = root/'run_app.py' if (root/'run_app.py').exists() else root/'run_release.py'
    source = entry.read_text(encoding='utf-8')
    assert source.index('main(sys.argv[2:])') < source.index('from linda_desktop.__main__ import main')
