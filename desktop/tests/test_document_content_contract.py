"""Document imports must preserve body text and apply decoded byte bounds."""
import io

import pytest

from linda_desktop import documents, document_worker


@pytest.mark.parametrize('encoding,text', [
    ('utf-8', 'a' * 65),
    ('cp1251', '\u0416' * 33),
    ('utf-16', '\u0416' * 33),
])
def test_plain_text_decoded_utf8_limit(encoding, text, monkeypatch):
    monkeypatch.setattr(documents, 'MAX_TEXT', 64)
    with pytest.raises(ValueError, match='Extracted text exceeds'):
        documents.extract('input.txt', text.encode(encoding))


@pytest.mark.parametrize('encoding,text', [
    ('utf-8', 'a' * 64),
    ('cp1251', '\u0416' * 32),
    ('utf-16', '\u0416' * 32),
])
def test_plain_text_at_decoded_limit_is_unchanged(encoding, text, monkeypatch):
    monkeypatch.setattr(documents, 'MAX_TEXT', 64)
    assert documents.extract('input.md', text.encode(encoding)) == text


def _bytes(document):
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _table_document():
    import docx
    document = docx.Document()
    document.add_paragraph('Before')
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Left'
    table.cell(0, 1).text = 'Right'
    merged = table.cell(1, 0).merge(table.cell(1, 1))
    merged.text = 'Merged'
    nested = merged.add_table(rows=1, cols=2)
    nested.cell(0, 0).text = 'Nested left'
    nested.cell(0, 1).text = 'Nested right'
    merged.add_paragraph('After nested')
    document.add_paragraph('After')
    return document


def test_docx_preserves_body_table_order_nested_and_merged_cells():
    text = document_worker.parse('input.docx', _bytes(_table_document()))
    assert text == '\n\n'.join([
        'Before', 'Left', 'Right', 'Merged', 'Nested left',
        'Nested right', 'After nested', 'After',
    ])


def test_docx_worker_preserves_tables_in_isolated_import(tmp_path, monkeypatch):
    monkeypatch.setenv('LINDA_HOME', str(tmp_path / 'document-profile'))
    assert 'Nested right' in documents.extract('input.docx', _bytes(_table_document()))


def test_docx_vertical_merge_is_not_duplicated_and_equal_cells_are_kept():
    import docx
    document = docx.Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(1, 0)).text = 'Merged'
    table.cell(0, 1).text = 'Repeated'
    table.cell(1, 1).text = 'Repeated'
    assert document_worker.parse('input.docx', _bytes(document)) == 'Merged\n\nRepeated\n\nRepeated'


def test_docx_table_text_counts_toward_output_limit(monkeypatch):
    import docx
    document = docx.Document()
    document.add_table(rows=1, cols=1).cell(0, 0).text = 'a' * 65
    monkeypatch.setattr(document_worker, 'MAX_TEXT', 64)
    with pytest.raises(ValueError, match='Extracted text exceeds'):
        document_worker.parse('input.docx', _bytes(document))


def test_docx_exact_output_limit_and_no_partial_text(monkeypatch):
    import docx
    document = docx.Document()
    document.add_paragraph('a' * 31)
    document.add_paragraph('b' * 31)
    monkeypatch.setattr(document_worker, 'MAX_TEXT', 64)
    assert document_worker.parse('input.docx', _bytes(document)) == 'a' * 31 + '\n\n' + 'b' * 31
    document.add_paragraph('tail')
    with pytest.raises(ValueError, match='Extracted text exceeds'):
        document_worker.parse('input.docx', _bytes(document))
